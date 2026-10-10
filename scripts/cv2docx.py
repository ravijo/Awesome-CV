#!/usr/bin/env python3
"""Generate a Word (.docx) version of an Awesome-CV document.

The DOCX is built from the same sources as the PDF: the top-level .tex file,
the section files it \\input's, and the .bbl file biber writes for the
publication lists. Layout, fonts, sizes and colours follow awesome-cv.cls, and
the fonts are embedded so the document looks the same on machines without them.
Japanese CVs (class option "ja") use a Word house style instead (see JP_STYLE).

Usage:
    python3 scripts/cv2docx.py examples/cv.tex [-o examples/cv.docx] [--bbl examples/cv.bbl]

Run it after the PDF build (it needs the .bbl) and before temporary files are
cleaned. Requires python-docx and fontTools, and TeX Live's kpsewhich to locate
the fonts.
"""

import argparse
import datetime
import hashlib
import io
import os
import re
import subprocess
import sys
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, replace

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

# -----------------------------------------------------------------------------
# Class configuration (mirrors awesome-cv.cls)
# -----------------------------------------------------------------------------
AWESOME_COLORS = {
    "awesome-emerald": "00A388", "awesome-skyblue": "0395DE", "awesome-red": "DC3522",
    "awesome-pink": "EF4089", "awesome-orange": "FF6138", "awesome-nephritis": "27AE60",
    "awesome-concrete": "95A5A6", "awesome-darknight": "131A28",
}
DEFAULT_COLORS = {
    "darktext": "414141", "text": "333333", "graytext": "5D5D5D", "lighttext": "999999",
    "sectiondivider": "5D5D5D", "lightdivider": "CECECE", "awesome": AWESOME_COLORS["awesome-red"],
}
COLORS = dict(DEFAULT_COLORS)  # per document; reset by Context

# Font roles -> OpenType files (regular, bold, italic, bold italic). A "light"
# role switches to its bold counterpart when \bfseries is applied, like LaTeX.
FONT_FILES = {
    "sans": ("SourceSans3-Regular.otf", "SourceSans3-Bold.otf", "SourceSans3-RegularIt.otf", "SourceSans3-BoldIt.otf"),
    "sanslight": ("SourceSans3-Light.otf", None, "SourceSans3-LightIt.otf", None),
    "head": ("Roboto-Regular.otf", "Roboto-Bold.otf", "Roboto-Italic.otf", "Roboto-BoldItalic.otf"),
    "headlight": ("Roboto-Light.otf", None, "Roboto-LightItalic.otf", None),
    "mono": ("lmmonolt10-regular.otf", None, "lmmonolt10-oblique.otf", None),  # \url in light text
    "cjk": ("HaranoAjiGothic-Regular.otf", "HaranoAjiGothic-Bold.otf", None, None),
    "fasolid": ("FontAwesome6Free-Solid-900.otf", None, None, None),
    "fabrands": ("FontAwesome6Brands-Regular-400.otf", None, None, None),
}
BOLD_OF = {"sanslight": "sans", "headlight": "head"}

# Japanese CVs (class option "ja") follow the house style of the reference Word files:
# standard Windows Japanese fonts (not embedded), black text and headings, no blue
# (the accent colour is black), and greys kept for job titles, dates and labels.
JP_STYLE = {
    # (Latin name, Japanese name) as Japanese Word writes them; the font table links the two
    "font": ("MS PGothic", "ＭＳ Ｐゴシック"),   # all text
    "name_font": ("MS Gothic", "ＭＳ ゴシック"),  # the name in the header
    "name_size": 24, "social_size": 6.5,
    "name_sep": "・",  # between all parts of the name
    # black text, and no blue: the accent colour (position line, locations, icons) is black too
    "colors": {"text": "000000", "awesome": "000000"},
}
# Font table entries for the Japanese system fonts (charset 80 = Shift-JIS)
JP_FONT_TABLE = {
    "ＭＳ ゴシック": ("MS Gothic", "modern", "fixed"),
    "ＭＳ Ｐゴシック": ("MS PGothic", "swiss", "variable"),
}
JP_MODE = False  # per document; set by Context

# Vertical layout, in points, calibrated against the PDF. "line" values are
# exact line heights; the rest are gaps, mostly added as space-after on the
# previous block (Builder.gap) so they vanish at page breaks like TeX glue.
LAYOUT = {
    # Header lines are as tight as possible without Word clipping the glyph tops
    "name_line": 30.8, "position_line": 8.5, "address_line": 11.2, "social_line": 12.2,
    "after_header": 4.0,
    # Word sets the first row of a page lower than TeX's \topskip does: pull the
    # top margin up by this much and give page 1 the space back before the header
    "page_top_shift": 3.0,
    "section_line": 21.0, "section_before": 6.0, "section_after": 2.0,
    "subsection_line": 16.0, "subsection_before": 1.5, "subsection_after": 0.0,
    "entry_title_line": 15.6, "entry_position_line": 12.9, "entry_gap": 3.8,
    "item_line": 10.9, "items_before": 2.0, "first_entry_gap": 3.0,
    "honor_line": 13.5, "honors_before": 2.5, "honors_after": 6.6,
    "entries_after": 9.7,
    "skill_line": 13.5, "skills_before": 1.9, "skills_after": 6.3,
    "paragraph_line": 12.0,
    "bib_title_line": 10.9, "bib_author_line": 10.7, "bib_gap": 2.95, "bib_after": 9.0, "bib_first_gap": 5.8,
    "bib_venue_line": 8.7,   # article: journal line (a paragraph)
    "bib_article_gap": 4.4,  # articles are followed by a little more space
    "bib_row_line": 9.0,     # rows of the venue/location and doi/year tables
    "bib_article_row_line": 9.9,
}
EX = 5.35       # 1ex of 11pt Source Sans 3 (x-height 486/1000)
SMALLCAPS = 0.85  # Source Sans 3 small-cap letters relative to capitals
EM = 11.0       # 1em of the 11pt body font
BIBHANG = 10.9  # biblatex \bibhang (1em)

# LaTeX accent commands -> Unicode combining marks (\'e -> é, \c{c} -> ç, ...)
ACCENTS = {
    "'": "\u0301", "`": "\u0300", "^": "\u0302", '"': "\u0308", "~": "\u0303", "=": "\u0304",
    ".": "\u0307", "c": "\u0327", "v": "\u030c", "u": "\u0306", "H": "\u030b", "k": "\u0328",
    "r": "\u030a", "d": "\u0323", "b": "\u0331",
}

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

# -----------------------------------------------------------------------------
# Fonts: locate, measure, convert to TrueType and embed
# -----------------------------------------------------------------------------
from fontTools.ttLib import TTFont, newTable  # noqa: E402
from fontTools.pens.cu2quPen import Cu2QuPen  # noqa: E402
from fontTools.pens.ttGlyphPen import TTGlyphPen  # noqa: E402

_font_cache = {}
_font_path_cache = {}


def font_path(filename):
    if filename not in _font_path_cache:
        try:
            p = subprocess.run(["kpsewhich", filename], capture_output=True, text=True).stdout.strip()
        except FileNotFoundError:
            p = ""
        _font_path_cache[filename] = p or None
    return _font_path_cache[filename]


def load_font(filename):
    if filename not in _font_cache:
        p = font_path(filename)
        _font_cache[filename] = TTFont(p, lazy=True) if p else None
    return _font_cache[filename]


def family_name(filename):
    f = load_font(filename)
    return f["name"].getDebugName(1) if f else filename.rsplit(".", 1)[0]


def resolve_font(role, bold, italic):
    """Map a role + bold/italic request to (role actually used, font file)."""
    if bold and role in BOLD_OF:
        role = BOLD_OF[role]
    files = FONT_FILES[role]
    f = files[(1 if bold else 0) + (2 if italic else 0)]
    if f is None:  # closest face that exists; Word synthesises the rest
        f = files[2] if italic and files[2] else files[0]
    return role, f


def text_width(text, role, size, bold=False, italic=False, smallcaps=False):
    """Advance width of text in points (no kerning)."""
    _, f = resolve_font(role, bold, italic)
    if JP_MODE and role not in ("fasolid", "fabrands"):
        f = FONT_FILES["cjk"][1 if bold else 0]  # stand-in for MS PGothic, which is not available here
    font = load_font(f)
    if font is None:
        return len(text) * size * 0.5
    cmap = font.getBestCmap()
    hmtx = font["hmtx"]
    upem = font["head"].unitsPerEm
    cjk = None
    w = 0
    for ch in text:
        s = size
        if smallcaps and ch.islower():
            ch, s = ch.upper(), size * SMALLCAPS
        g = cmap.get(ord(ch))
        if g:
            w += hmtx[g][0] * s / upem
            continue
        if cjk is None:  # Word falls back to the East Asian font for these
            cjk = load_font(FONT_FILES["cjk"][1 if bold else 0]) or False
        g = cjk.getBestCmap().get(ord(ch)) if cjk else None
        if g:
            w += cjk["hmtx"][g][0] * s / cjk["head"].unitsPerEm
        else:
            w += s * 0.5
    return w


def has_glyph(role, ch):
    font = load_font(FONT_FILES[role][0])
    return font is not None and ord(ch) in font.getBestCmap()


def otf_to_ttf(src, dst):
    """Convert a CFF-flavoured OpenType font to TrueType outlines (Word only embeds those).
    The source's timestamp is kept, so converting the same font always gives the same bytes."""
    font = TTFont(src, recalcTimestamp=False)
    if "glyf" in font:
        font.save(dst)
        return
    order = font.getGlyphOrder()
    gs = font.getGlyphSet()
    glyphs = {}
    for name in order:
        pen = TTGlyphPen(gs)
        gs[name].draw(Cu2QuPen(pen, 1.0, reverse_direction=True))
        glyphs[name] = pen.glyph()
    font["loca"] = newTable("loca")
    font["glyf"] = glyf = newTable("glyf")
    glyf.glyphOrder = order
    glyf.glyphs = glyphs
    del font["CFF "]
    if "VORG" in font:
        del font["VORG"]
    glyf.compile(font)
    hmtx = font["hmtx"]
    for name, g in glyf.glyphs.items():
        if hasattr(g, "xMin"):
            hmtx[name] = (hmtx[name][0], g.xMin)
    font["maxp"] = maxp = newTable("maxp")
    maxp.tableVersion = 0x00010000
    for attr in ("maxZones", "maxTwilightPoints", "maxStorage", "maxFunctionDefs",
                 "maxInstructionDefs", "maxStackElements", "maxSizeOfInstructions"):
        setattr(maxp, attr, 0)
    maxp.maxZones = 1
    maxp.maxComponentElements = max((len(getattr(g, "components", []) or []) for g in glyphs.values()), default=0)
    maxp.compile(font)
    post = font["post"]
    post.formatType = 2.0
    post.extraNames = []
    post.mapping = {}
    post.glyphOrder = order
    try:
        post.compile(font)
    except OverflowError:
        post.formatType = 3.0
    # The outlines are unhinted: tell renderers (Windows/Word in particular) to use
    # smoothing at every size rather than grid-fitting, via a gasp table
    gasp = newTable("gasp")
    gasp.version = 1
    gasp.gaspRange = {0xFFFF: 0x000A}  # GASP_DOGRAY | GASP_SYMMETRIC_SMOOTHING
    font["gasp"] = gasp
    font.sfntVersion = "\000\001\000\000"
    font.save(dst)


CONVERT_VERSION = 2  # bump when otf_to_ttf changes, so cached fonts are rebuilt


def ttf_bytes(filename, cache_dir):
    """TrueType version of a font, cached in cache_dir."""
    src = font_path(filename)
    if not src:
        return None
    st = os.stat(src)
    key = hashlib.sha1(f"{CONVERT_VERSION}:{src}:{st.st_size}:{st.st_mtime}".encode()).hexdigest()[:12]
    dst = os.path.join(cache_dir, f"{os.path.splitext(filename)[0]}-{key}.ttf")
    if not os.path.exists(dst):
        os.makedirs(cache_dir, exist_ok=True)
        tmp = f"{dst}.{os.getpid()}.tmp"
        otf_to_ttf(src, tmp)
        os.replace(tmp, dst)
    with open(dst, "rb") as fh:
        return fh.read()


def obfuscate(data, key_guid):
    """ECMA-376 font obfuscation: XOR the first 32 bytes with the reversed GUID bytes."""
    key = bytes.fromhex(key_guid.strip("{}").replace("-", ""))[::-1]
    data = bytearray(data)
    for i in range(32):
        data[i] ^= key[i % 16]
    return bytes(data)


# -----------------------------------------------------------------------------
# Minimal LaTeX reading
# -----------------------------------------------------------------------------
def strip_comments(text):
    return re.sub(r"(?<!\\)%.*", "", text)


def skip_ws(s, i):
    while i < len(s) and s[i] in " \t\r\n":
        i += 1
    return i


def read_group(s, i):
    """Read a {...} argument (or a single token) starting at i; returns (content, next index)."""
    i = skip_ws(s, i)
    if i >= len(s):
        return "", i
    if s[i] != "{":
        if s[i] == "\\":
            m = re.match(r"\\([A-Za-z@]+|.)", s[i:])
            return m.group(0), i + len(m.group(0))
        return s[i], i + 1
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def read_optional(s, i):
    """Read an optional [...] argument; returns (content or None, next index)."""
    j = skip_ws(s, i)
    if j >= len(s) or s[j] != "[":
        return None, i
    depth, k = 0, j
    while k < len(s):
        c = s[k]
        if c == "\\":
            k += 2
            continue
        if c in "{[":
            depth += 1
        elif c in "}]":
            depth -= 1
            if depth == 0 and c == "]":
                return s[j + 1:k], k + 1
        k += 1
    return s[j + 1:], len(s)


def find_env_end(s, i, env):
    """Index of the matching \\end{env} for an environment whose body starts at i."""
    depth = 1
    pat = re.compile(r"\\(begin|end)\{" + re.escape(env) + r"\}")
    for m in pat.finditer(s, i):
        depth += 1 if m.group(1) == "begin" else -1
        if depth == 0:
            return m.start(), m.end()
    return len(s), len(s)


def parse_length(text, default=0.0):
    """Convert a LaTeX length like 2.5cm / 4mm / 10pt / 1em to points."""
    if text is None:
        return default
    text = text.strip()
    m = re.fullmatch(r"(-?[\d.]+)\s*(cm|mm|pt|in|em|ex|bp)", text)
    if not m:
        return default
    v, u = float(m.group(1)), m.group(2)
    return v * {"cm": 28.4528, "mm": 2.84528, "pt": 1.0, "in": 72.27, "em": EM, "ex": EX, "bp": 1.00375}[u]


# -----------------------------------------------------------------------------
# Inline text -> formatted segments
# -----------------------------------------------------------------------------
@dataclass(frozen=True)
class Fmt:
    role: str = "sanslight"
    size: float = 9.0
    color: str = "text"
    bold: bool = False
    italic: bool = False
    smallcaps: bool = False
    link: str = None
    spacing: float = 0.0  # character spacing in points (negative = condensed)


@dataclass
class Seg:
    text: str
    fmt: Fmt
    kind: str = "text"  # text | br | field


class Inline:
    """Turns a LaTeX fragment into Seg runs."""

    def __init__(self, ctx):
        self.ctx = ctx
        self.icons = ctx.icons

    def __call__(self, s, fmt):
        segs = self._parse(s, fmt)
        return normalize_space(segs)

    def _parse(self, s, fmt):
        out, buf = [], []

        def flush():
            if buf:
                out.append(Seg("".join(buf), fmt))
                buf.clear()

        i = 0
        while i < len(s):
            c = s[i]
            if c == "\\":
                if i + 1 < len(s) and not s[i + 1].isalpha() and s[i + 1] != "@":
                    sym = s[i + 1]
                    i += 2
                    if sym in ACCENTS:
                        arg, i = read_group(s, i)
                        buf.append(self._accent(arg, ACCENTS[sym], fmt))
                    elif sym in "&%$#_{}":
                        buf.append(sym)
                    elif sym == "\\":
                        flush()
                        out.append(Seg("\n", fmt, "br"))
                        i = skip_ws(s, i)
                        if i < len(s) and s[i] == "*":
                            i += 1
                        # \\[2mm] is spacing; anything else in brackets is text
                        m = re.match(r"\s*\[\s*-?[\d.]+\s*(?:pt|mm|cm|in|em|ex|bp)\s*\]", s[i:])
                        if m:
                            i += m.end()
                    elif sym == ",":
                        buf.append("\u2009")
                    elif sym in " \n":
                        buf.append(" ")
                    continue
                m = re.match(r"[A-Za-z@]+", s[i + 1:])
                if not m:  # a lone backslash at the end of the text: ignore it
                    i += 1
                    continue
                name = m.group(0)
                i += 1 + len(m.group(0))
                j = i
                while j < len(s) and s[j] in " \t\n":
                    j += 1
                i = j
                flush()
                i = self._command(name, s, i, fmt, out, buf)
            elif c == "{":
                arg, i = read_group(s, i)
                flush()
                out.extend(self._parse(arg, fmt))
            elif c == "}":
                i += 1
            elif c == "$":
                j = s.find("$", i + 1)
                j = len(s) if j < 0 else j
                flush()
                out.extend(self._parse(s[i + 1:j], fmt))
                i = j + 1
            elif c == "~":
                buf.append("\u00a0")
                i += 1
            elif c == "-":
                if s.startswith("---", i):
                    buf.append("\u2014")
                    i += 3
                elif s.startswith("--", i):
                    buf.append("\u2013")
                    i += 2
                else:
                    buf.append("-")
                    i += 1
            elif c == "`":
                if s.startswith("``", i):
                    buf.append("\u201c")
                    i += 2
                else:
                    buf.append("\u2018")
                    i += 1
            elif c == "'":
                if s.startswith("''", i):
                    buf.append("\u201d")
                    i += 2
                else:
                    buf.append("\u2019")
                    i += 1
            elif c in " \t\r\n":
                buf.append(" ")
                i += 1
            else:
                buf.append(c)
                i += 1
        flush()
        return out

    def _command(self, name, s, i, fmt, out, buf):
        styled = {
            "textbf": dict(bold=True), "textit": dict(italic=True), "emph": dict(italic=True),
            "textsl": dict(italic=True), "textsc": dict(smallcaps=True), "texttt": dict(role="mono"),
            "textup": {}, "textnormal": {}, "textrm": {}, "textsf": {}, "textmd": dict(bold=False),
            "mbox": {}, "makebox": {}, "underline": {},
        }
        symbols = {
            "textbullet": "\u2022", "bullet": "\u2022", "LaTeX": "LaTeX", "TeX": "TeX", "O": "\u00d8",
            "o": "\u00f8", "textcopyright": "\u00a9", "textregistered": "\u00ae", "texttrademark": "\u2122",
            "cdotp": "\u00b7", "cdot": "\u00b7", "times": "\u00d7", "textbar": "|", "ldots": "\u2026",
            "dots": "\u2026", "textellipsis": "\u2026", "space": " ", "textbackslash": "\\",
            "textendash": "\u2013", "textemdash": "\u2014", "pm": "\u00b1", "sim": "~", "AA": "\u00c5",
            # common \item[...] labels
            "star": "\u22c6", "ast": "\u2217", "circ": "\u2218", "diamond": "\u22c4", "checkmark": "\u2713",
            "rightarrow": "\u2192", "to": "\u2192", "textasteriskcentered": "\u2217", "textopenbullet": "\u25e6",
        }
        if name in ACCENTS:  # letter accents such as \c{c} or \v{s}
            arg, i = read_group(s, i)
            buf.append(self._accent(arg, ACCENTS[name], fmt))
        elif name in ("i", "j"):  # dotless i/j, used under accents
            buf.append("\u0131" if name == "i" else "\u0237")
        elif name in styled:
            if name == "makebox":
                _, i = read_optional(s, i)
                _, i = read_optional(s, i)
            arg, i = read_group(s, i)
            out.extend(self._parse(arg, replace(fmt, **styled[name])))
        elif name in symbols:
            buf.append(symbols[name])
            if s[i:i + 2] == "{}":
                i += 2
        elif name in ("enskip", "quad", "qquad"):
            buf.append(self.ctx.space(fmt, {"enskip": 0.5, "quad": 1.0, "qquad": 2.0}[name]))
        elif name == "today":
            buf.append(self.ctx.today())
        elif name == "href":
            url, i = read_group(s, i)
            txt, i = read_group(s, i)
            out.extend(self._parse(txt, replace(fmt, link=unescape_url(url))))
        elif name == "url":
            url, i = read_group(s, i)
            out.append(Seg(url, replace(fmt, role="mono", link=url)))
        elif name == "webref":
            txt, i = read_group(s, i)
            url, i = read_group(s, i)
            lf = replace(fmt, link=unescape_url(url))
            out.extend(self._parse(txt, lf))
            out.append(Seg("\u00a0", lf))
            out.append(self.icon("faArrowUpRightFromSquare", lf))
        elif name.startswith("fa") and name in self.icons:
            out.append(self.icon(name, fmt))
        elif name == "thepage":
            out.append(Seg("PAGE", fmt, "field"))
        elif name == "pageref":
            ref, i = read_group(s, i)
            if ref == "LastPage":
                out.append(Seg("NUMPAGES", fmt, "field"))
        elif name in ("hspace", "vspace"):
            if i < len(s) and s[i] == "*":
                i += 1
            _, i = read_group(s, i)
        elif name in ("newline", "linebreak"):
            out.append(Seg("\n", fmt, "br"))
        elif name in ("begin", "end"):  # stray environment markers carry no text
            _, i = read_group(s, i)
            if name == "begin":
                while read_optional(s, i)[0] is not None:
                    _, i = read_optional(s, i)
        elif name == "refsectionbibcount":
            buf.append(str(self.ctx.bibcount))
        else:
            # Unknown command: keep the text of a following argument, if any
            j = skip_ws(s, i)
            if j < len(s) and s[j] == "{":
                arg, i = read_group(s, j)
                out.extend(self._parse(arg, fmt))
        return i

    def _accent(self, arg, mark, fmt):
        """Put a combining accent on the first letter of arg and compose it (e + ´ -> é)."""
        text = plain(self._parse(arg, fmt))
        if not text:
            return ""
        base = {"\u0131": "i", "\u0237": "j"}.get(text[0], text[0])  # \'\i -> í
        return unicodedata.normalize("NFC", base + mark + text[1:])

    def icon(self, name, fmt):
        role, ch = self.icons[name]
        return Seg(ch, replace(fmt, role=role, bold=False, italic=False, smallcaps=False))


def normalize_space(segs):
    """Collapse whitespace like TeX: runs of spaces become one, none at the ends."""
    out = []
    prev_space = True
    for seg in segs:
        if seg.kind != "text":
            out.append(seg)
            prev_space = seg.kind == "br"
            continue
        t = re.sub(r"[ \t\r\n]+", " ", seg.text)
        if prev_space and t.startswith(" "):
            t = t[1:]
        if t:
            out.append(Seg(t, seg.fmt))
            prev_space = t.endswith(" ")
    while out and out[-1].kind == "text" and out[-1].text.endswith(" "):
        t = out[-1].text.rstrip(" ")
        if t:
            out[-1] = Seg(t, out[-1].fmt)
            break
        out.pop()
    return out


def line_count(words, avail, shrink_per_space):
    """Greedy line breaking; words are (width, space after) pairs. A line may
    exceed avail by shrink_per_space for every space on it."""
    lines, cur, spaces = 1, 0.0, 0
    for w, sp in words:
        if cur and cur + w > avail + spaces * shrink_per_space:
            lines += 1
            cur, spaces = 0.0, 0
        cur += w + sp
        spaces += 1
    return lines


def condense(segs, avail):
    """Word only stretches spaces in justified text, while TeX may also shrink them
    (by a third of a space). Where that lets TeX set a paragraph in fewer lines,
    tighten the character spacing just enough for Word to do the same."""
    if any(s.kind != "text" for s in segs):
        return segs
    chars = []  # (char, width, fmt)
    for s in segs:
        for ch in s.text:
            chars.append((ch, text_width(ch, s.fmt.role, s.fmt.size, s.fmt.bold, s.fmt.italic, s.fmt.smallcaps), s.fmt))
    if not chars:
        return segs
    def words(c, kern=0.995):  # pair kerning: ~0.5% narrower in Word, ~1% in TeX
        out, w = [], 0.0
        for ch, cw, _ in chars:
            if ch == " ":
                out.append((w, (cw - c) * kern))
                w = 0.0
            else:
                w += (cw - c) * kern
        out.append((w, 0.0))
        return out

    # TeX: a line may exceed the measure by a third of each of its spaces
    third = text_width(" ", "sanslight", 9) / 3
    target = line_count(words(0.0, 0.99), avail, third)
    if target >= line_count(words(0.0), avail, 0.0):
        return segs
    for step in range(1, 7):  # up to 0.3pt per character
        c = step * 0.05
        if line_count(words(c), avail, 0.0) <= target:
            return [Seg(s.text, replace(s.fmt, spacing=-c), s.kind) for s in segs]
    return segs


def unescape_url(url):
    """Drop LaTeX escapes from a URL, as hyperref does (a\\_b\\%20c\\#top -> a_b%20c#top)."""
    return re.sub(r"\\([_%#&~$])", r"\1", url.strip())


def plain(segs):
    return "".join(s.text for s in segs if s.kind == "text")


# -----------------------------------------------------------------------------
# .bbl reading
# -----------------------------------------------------------------------------
def parse_bbl(path):
    """Return {refsection number: [entry dict, ...]} in biber's sort order."""
    sections = {}
    if not path or not os.path.exists(path):
        return sections
    text = open(path, encoding="utf-8").read()
    for sm in re.finditer(r"\\refsection\{(\d+)\}(.*?)\\endrefsection", text, re.S):
        entries, seen = [], set()
        for em in re.finditer(r"\\entry\{([^}]*)\}\{([^}]*)\}.*?\n(.*?)\\endentry", sm.group(2), re.S):
            key, etype, body = em.groups()
            if key in seen:
                continue
            seen.add(key)
            e = {"key": key, "type": etype, "fields": {}, "lists": {}, "names": {}, "verbs": {}}
            for fm in re.finditer(r"\\field\{(\w+)\}", body):
                v, _ = read_group(body, fm.end())
                e["fields"][fm.group(1)] = v
            for lm in re.finditer(r"\\list\{(\w+)\}\{\d+\}\{%", body):
                v, _ = read_group(body, lm.end() - 2)
                e["lists"][lm.group(1)] = re.findall(r"\{((?:[^{}]|\{[^{}]*\})*)\}%", v)
            for nm in re.finditer(r"\\name\{(\w+)\}\{\d+\}\{[^}]*\}\{%", body):
                v, _ = read_group(body, nm.end() - 2)
                # each person starts "{{hash=…" (or "{{un=0,…,hash=…" with uniquename options)
                e["names"][nm.group(1)] = [name_parts(p) for p in re.split(r"\{\{(?:\w+=[^,{}]*,)*hash=", v)[1:]]
            for vm in re.finditer(r"\\verb\{(\w+)\}\s*\\verb (.*?)\n\s*\\endverb", body, re.S):
                e["verbs"][vm.group(1)] = vm.group(2).strip()
            entries.append(e)
        sections[int(sm.group(1))] = entries
    return sections


def name_parts(person):
    """One person from a .bbl name list -> {"given", "prefix", "family", "suffix"}."""
    parts = {}
    for m in re.finditer(r"\b(given|prefix|family|suffix)=", person):
        parts[m.group(1)], _ = read_group(person, m.end())
    return parts


def format_name(parts):
    """Given, von-part, family, Jr-part, as the class's author name format prints them."""
    name = " ".join(x for x in (parts.get("given"), parts.get("prefix"), parts.get("family")) if x)
    return name + (" " + parts["suffix"] if parts.get("suffix") else "")


def bib_tex(v):
    """Field text from the .bbl as LaTeX for the inline renderer."""
    v = v.replace("\\bibrangedash", "--").replace("\\bibnamedelima", " ").replace("\\bibnamedelimb", " ")
    v = v.replace("\\bibnamedelimc", " ").replace("\\bibnamedelimd", " ").replace("\\bibnamedelimi", " ")
    v = re.sub(r"\\bibrangessep\s*", ", ", v)
    return v


# -----------------------------------------------------------------------------
# DOCX building
# -----------------------------------------------------------------------------
class Builder:
    def __init__(self, ctx):
        self.ctx = ctx
        self.doc = Document()
        self.used_fonts = set()  # (family, slot, file)
        self.first_in_section = False
        self._bullet_num = None
        self._setup_document()

    # --- document setup ---------------------------------------------------
    def _setup_document(self):
        ctx = self.ctx
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Pt(ctx.paper[0]), Pt(ctx.paper[1])
        sec.left_margin, sec.right_margin = Pt(ctx.margins["left"]), Pt(ctx.margins["right"])
        sec.top_margin = Pt(ctx.margins["top"] - LAYOUT["page_top_shift"])
        sec.bottom_margin = Pt(ctx.margins["bottom"])
        sec.header_distance = Pt(ctx.margins["top"] / 2)
        sec.footer_distance = Pt(max(ctx.margins["bottom"] - ctx.margins["footskip"] - 2.1, 6))
        normal = self.doc.styles["Normal"]
        pf = normal.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.widow_control = False
        self._set_run_fonts(normal.element.get_or_add_rPr(), "sanslight", False, False)
        normal.font.size = Pt(9)
        lang = "ja-JP" if ctx.japanese else "en-US"
        rpr = normal.element.get_or_add_rPr()
        el = OxmlElement("w:lang")
        el.set(qn("w:val"), "en-US")
        el.set(qn("w:eastAsia"), lang)
        rpr.append(el)
        settings = self.doc.settings.element
        hy = OxmlElement("w:autoHyphenation")
        settings.append(hy)
        # python-docx's template declares Word 2010 (14), which makes Word open the file
        # in "Compatibility Mode"; 15 is the current layout mode
        for cs in settings.iter(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                cs.set(qn("w:val"), "15")
        props = self.doc.core_properties
        props.author = ctx.pdfauthor
        props.title = ctx.pdftitle
        props.subject = ctx.pdftitle
        props.keywords = ""
        props.comments = ""

    @property
    def text_width(self):
        return self.ctx.paper[0] - self.ctx.margins["left"] - self.ctx.margins["right"]

    # --- runs --------------------------------------------------------------
    def _set_run_fonts(self, rpr, role, bold, italic):
        if self.ctx.japanese and role not in ("fasolid", "fabrands"):
            latin, japanese = JP_STYLE["name_font"] if role == "name" else JP_STYLE["font"]
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(attr), latin)
            rfonts.set(qn("w:eastAsia"), japanese)
            return latin
        role, f = resolve_font(role, bold, italic)
        fam = family_name(f)
        slot = (1 if bold else 0) + (2 if italic else 0)
        files = FONT_FILES[role]
        if files[slot] is None:
            slot = 2 if italic and files[2] else 0
        self.used_fonts.add((fam, slot, files[slot]))
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), fam)
        rfonts.set(qn("w:eastAsia"), fam)
        return fam

    def add_runs(self, par, segs):
        for seg in segs:
            if seg.kind == "br":
                par.add_run().add_break()
                continue
            if seg.kind == "field":
                self._field(par, seg)
                continue
            if seg.fmt.link:
                self._hyperlink(par, seg)
            else:
                self._run(par, seg)

    def _style_run(self, run, fmt):
        f = fmt
        rpr = run._r.get_or_add_rPr()
        self._set_run_fonts(rpr, f.role, f.bold, f.italic)
        run.font.size = Pt(f.size)
        run.font.bold = bool(f.bold)
        run.font.italic = bool(f.italic)
        if f.smallcaps and JP_MODE:
            run.font.small_caps = True
        color = COLORS.get(f.color, f.color)
        run.font.color.rgb = RGBColor.from_string(color)
        if f.spacing:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(int(round(f.spacing * 20))))
            rpr.append(sp)
        kern = OxmlElement("w:kern")  # LaTeX kerns these fonts; Word only does when asked
        kern.set(qn("w:val"), "2")
        rpr.append(kern)

    @staticmethod
    def _pieces(seg):
        """Split a small-caps segment into capitals and reduced capitals, as Word's
        synthetic small caps are much smaller than the fonts' real ones."""
        if not seg.fmt.smallcaps or JP_MODE:
            return [(seg.text, seg.fmt)]
        out = []
        for m in re.finditer(r"[^\W\d_]+|[\W\d_]+", seg.text):
            chunk = m.group(0)
            for part in re.findall(r"[a-z\u00df-\u00ff]+|[^a-z\u00df-\u00ff]+", chunk):
                if part[0].islower():
                    out.append((part.upper(), replace(seg.fmt, size=round(seg.fmt.size * SMALLCAPS * 2) / 2)))
                else:
                    out.append((part, seg.fmt))
        return out

    def _run(self, par, seg):
        for text, fmt in self._pieces(seg):
            run = par.add_run(text)
            self._style_run(run, fmt)

    def _hyperlink(self, par, seg):
        rid = par.part.relate_to(seg.fmt.link, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        for text, fmt in self._pieces(seg):
            tmp = par.add_run(text)
            self._style_run(tmp, fmt)
            r = tmp._r
            r.getparent().remove(r)
            h.append(r)
        par._p.append(h)

    def _field(self, par, seg):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), seg.text)
        tmp = par.add_run("1")
        self._style_run(tmp, seg.fmt)
        r = tmp._r
        r.getparent().remove(r)
        fld.append(r)
        par._p.append(fld)

    # --- paragraphs ----------------------------------------------------------
    def para(self, container=None, line=None, align=None, before=0.0, after=0.0, keep=False, style=None,
             left=0.0, hanging=0.0):
        container = container or self.doc
        p = container.add_paragraph(style=style)
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if line:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(line)
        if align:
            p.alignment = align
        if left or hanging:
            pf.left_indent = Pt(left)
            pf.first_line_indent = Pt(-hanging)
        pf.keep_with_next = keep
        pf.widow_control = False
        return p

    def gap(self, height):
        """Vertical space after the previous block. Added as space-after so that,
        like TeX glue, it disappears at a page break instead of starting the page."""
        if height <= 0:
            return
        body = self.doc.element.body
        blocks = [el for el in body if el.tag in (qn("w:p"), qn("w:tbl"))]
        if not blocks:
            self.spacer(height)
            return
        last = blocks[-1]
        paras = [last] if last.tag == qn("w:p") else last.findall(qn("w:tr"))[-1].findall(".//" + qn("w:p"))
        for p in paras:
            pPr = p.get_or_add_pPr()
            sp = pPr.find(qn("w:spacing"))
            if sp is None:
                sp = OxmlElement("w:spacing")
                pPr.append(sp)
            cur = int(sp.get(qn("w:after"), "0"))
            sp.set(qn("w:after"), str(cur + int(round(height * 20))))

    def spacer(self, height, keep=True):
        if height <= 0:
            return
        p = self.para(line=max(height, 0.7), keep=keep)
        r = p.add_run("")
        r.font.size = Pt(1)
        pr = p._p.get_or_add_pPr()
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "2")
        rpr.append(sz)
        pr.append(rpr)

    # --- tables ----------------------------------------------------------------
    def table(self, widths, rows):
        t = self.doc.add_table(rows=0, cols=len(widths))
        t.autofit = False
        tbl = t._tbl
        tblPr = tbl.tblPr
        for tag, attrs in (("w:tblW", {"w:w": str(int(sum(widths) * 20)), "w:type": "dxa"}),
                           ("w:tblInd", {"w:w": "0", "w:type": "dxa"}),
                           ("w:tblLayout", {"w:type": "fixed"})):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            old = tblPr.find(qn(tag))
            if old is not None:
                tblPr.remove(old)
            tblPr.append(el)
        mar = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), "0")
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblPr.append(mar)
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 20)))
        for r_i in range(rows):
            row = t.add_row()
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            for c, w in zip(row.cells, widths):
                c.width = Pt(w)
                tcPr = c._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(int(w * 20)))
                tcW.set(qn("w:type"), "dxa")
        return t

    def cell_para(self, cell, segs, line, align=None, keep=True, left_pad=0.0, right_pad=0.0, valign=None):
        p = cell.paragraphs[0]
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
        pf.keep_with_next = keep
        pf.widow_control = False
        if left_pad:
            pf.left_indent = Pt(left_pad)
        if right_pad:
            pf.right_indent = Pt(right_pad)
        if align:
            p.alignment = align
        if valign:
            cell.vertical_alignment = valign
        self.add_runs(p, segs)
        return p

    # --- bullets -----------------------------------------------------------------
    def bullet_num(self):
        if self._bullet_num is not None:
            return self._bullet_num
        numbering = self.doc.part.numbering_part.element
        abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        aid, nid = max(abs_ids, default=0) + 1, max(num_ids, default=0) + 1
        indent = int(2 * 9 * 0.486 * 20 + 0.5)  # 2ex of the 9pt item text, in twips
        fam_runs = JP_STYLE["font"][0] if JP_MODE else family_name(FONT_FILES["sans"][0])
        fam_ea = JP_STYLE["font"][1] if JP_MODE else fam_runs
        xml = (
            f'<w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{aid}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\u2022"/>'
            '<w:lvlJc w:val="left"/>'
            f'<w:pPr><w:tabs><w:tab w:val="num" w:pos="{indent}"/></w:tabs><w:ind w:left="{indent}" w:hanging="{indent}"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{fam_runs}" w:hAnsi="{fam_runs}" w:cs="{fam_runs}" w:eastAsia="{fam_ea}"/>'
            f'<w:color w:val="{COLORS["text"]}"/><w:sz w:val="18"/></w:rPr>'
            '</w:lvl></w:abstractNum>')
        abs_el = etree.fromstring(xml)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abs_el)
        else:
            numbering.append(abs_el)
        num = etree.fromstring(
            f'<w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:numId="{nid}">'
            f'<w:abstractNumId w:val="{aid}"/></w:num>')
        numbering.append(num)
        if not JP_MODE:
            self.used_fonts.add((fam_runs, 0, FONT_FILES["sans"][0]))
        self._bullet_num = nid
        return nid

    def bullet(self, segs, keep, label=None):
        """A list item; with label (\\item[label]) the label replaces the bullet."""
        indent = 2 * 9 * 0.486
        if not JP_MODE:
            segs = condense(segs, self.text_width - indent)
        p = self.para(line=LAYOUT["item_line"], align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep=keep)
        if label is not None:
            # Hanging indent as for bullets; Word tabs from the label to the indent
            p.paragraph_format.left_indent = Pt(indent)
            p.paragraph_format.first_line_indent = Pt(-indent)
            self.add_runs(p, label + [Seg("\t", label[-1].fmt if label else Fmt())] + segs)
            return p
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(self.bullet_num()))
        numPr.append(il)
        numPr.append(ni)
        pPr.insert(0, numPr)
        self.add_runs(p, segs)
        return p

    # --- page break --------------------------------------------------------------
    def page_break(self):
        p = self.para(line=1)
        p.add_run().add_break(WD_BREAK.PAGE)

    # --- save with embedded fonts ------------------------------------------------
    def save(self, path, cache_dir, embed=True):
        doc = self.doc
        # Dates from the sources (not the clock), so unchanged sources give an identical file
        stamp = self.ctx.source_date()
        props = doc.core_properties
        props.created = props.modified = stamp
        props.last_modified_by = self.ctx.pdfauthor
        props.revision = 1
        for root in (doc.element, doc.styles.element, doc.settings.element, doc.part.numbering_part.element,
                     doc.sections[0].footer.part.element):
            sort_schema(root)
        buf = io.BytesIO()
        self.doc.save(buf)
        data = buf.getvalue()
        if JP_MODE:
            data = add_jp_font_table(data)
        if embed:
            data = embed_fonts(data, self.used_fonts, cache_dir)
        data = repack(data, stamp)
        with open(path, "wb") as fh:
            fh.write(data)


# Child order required by the OOXML schema (Word refuses files that break it)
SCHEMA_ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof "
           "snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd "
           "fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
           "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
             "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                "embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins "
                "alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop "
                "hideSpellingErrors hideGrammaticalErrors activeWritingStyle proofState formsDesign "
                "attachedTemplate linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
                "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection "
                "autoFormatOverride styleLockTheme styleLockQFSet defaultTabStop autoHyphenation "
                "consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength "
                "clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting "
                "bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing "
                "displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
                "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin "
                "doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne "
                "strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore savePreviewPicture "
                "doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText "
                "doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags "
                "alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars "
                "rsids mathPr attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats "
                "doNotAutoCompressPictures forceUpgrade captions readModeInkLockDown smartTagType "
                "schemaLibrary shapeDefaults doNotEmbedSmartTags decimalSymbol listSeparator",
}
SCHEMA_ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in SCHEMA_ORDER.items()}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def sort_schema(root):
    """Reorder children of property elements into schema order, merging duplicates
    (the last one wins). Unknown/extension elements keep their place at the end."""
    for name, order in SCHEMA_ORDER.items():
        nodes = [root] if etree.QName(root).localname == name else list(root.iter(f"{{{W_NS}}}{name}"))
        for node in nodes:
            kids = list(node)
            if len(kids) < 2:
                continue
            seen = {}
            for k in kids:
                q = etree.QName(k)
                if q.namespace == W_NS and q.localname in order:
                    if q.localname in seen and q.localname not in ("tab",):
                        node.remove(seen[q.localname])
                    seen[q.localname] = k
            kids = list(node)
            big = len(order) + 1
            kids.sort(key=lambda k: order.get(etree.QName(k).localname, big)
                      if etree.QName(k).namespace == W_NS else big)
            for k in kids:
                node.append(k)


def repack(docx_bytes, stamp):
    """Rewrite the package with every entry timestamped `stamp` (zip entries otherwise
    carry the current time, so identical content would still give a different file)."""
    date_time = max(stamp.timetuple()[:6], (1980, 1, 1, 0, 0, 0))
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            info = zipfile.ZipInfo(name, date_time=date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zout.writestr(info, zin.read(name))
    return out.getvalue()


def add_jp_font_table(docx_bytes):
    """Declare the Japanese system fonts in fontTable.xml under their Japanese names
    with the Latin names as altName, so Word resolves either spelling."""
    W = W_NS
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes))
    files = {n: zin.read(n) for n in zin.namelist()}
    table = etree.fromstring(files["word/fontTable.xml"])
    for jp, (latin, family, pitch) in JP_FONT_TABLE.items():
        for name in (jp, latin):
            node = table.find(f"{{{W}}}font[@{{{W}}}name='{name}']")
            if node is not None:
                table.remove(node)
        node = etree.SubElement(table, f"{{{W}}}font")
        node.set(f"{{{W}}}name", jp)
        etree.SubElement(node, f"{{{W}}}altName").set(f"{{{W}}}val", latin)
        for tag, val in (("charset", "80"), ("family", family), ("pitch", pitch)):
            etree.SubElement(node, f"{{{W}}}{tag}").set(f"{{{W}}}val", val)
    files["word/fontTable.xml"] = etree.tostring(table, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            zout.writestr(name, files[name])
    return out.getvalue()


def embed_fonts(docx_bytes, used, cache_dir):
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes))
    files = {n: zin.read(n) for n in zin.namelist()}
    font_table = etree.fromstring(files["word/fontTable.xml"])
    rels = etree.Element("{http://schemas.openxmlformats.org/package/2006/relationships}Relationships",
                         nsmap={None: "http://schemas.openxmlformats.org/package/2006/relationships"})
    slot_tag = {0: "embedRegular", 1: "embedBold", 2: "embedItalic", 3: "embedBoldItalic"}
    by_family = {}
    for fam, slot, f in sorted(used, key=lambda x: (x[0], x[1])):
        by_family.setdefault(fam, {})[slot] = f
    n = 0
    for fam, slots in by_family.items():
        node = font_table.find(f"{{{W}}}font[@{{{W}}}name='{fam}']")
        if node is None:
            node = etree.SubElement(font_table, f"{{{W}}}font")
            node.set(f"{{{W}}}name", fam)
            for tag, val in (("charset", "00"), ("family", "swiss"), ("pitch", "variable")):
                etree.SubElement(node, f"{{{W}}}{tag}").set(f"{{{W}}}val", val)
        for slot in sorted(slots):
            data = ttf_bytes(slots[slot], cache_dir)
            if data is None:
                continue
            n += 1
            # Key derived from the font itself: same font, same obfuscated bytes on every build
            key = "{" + str(uuid.uuid5(uuid.NAMESPACE_URL, "cv2docx:" + hashlib.sha1(data).hexdigest())).upper() + "}"
            name = f"font{n}.odttf"
            files[f"word/fonts/{name}"] = obfuscate(data, key)
            rid = f"rIdFont{n}"
            rel = etree.SubElement(rels, "Relationship")
            rel.set("Id", rid)
            rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
            rel.set("Target", f"fonts/{name}")
            emb = etree.SubElement(node, f"{{{W}}}{slot_tag[slot]}")
            emb.set(f"{{{R}}}id", rid)
            emb.set(f"{{{W}}}fontKey", key)
    if n == 0:
        return docx_bytes
    files["word/fontTable.xml"] = etree.tostring(font_table, xml_declaration=True, encoding="UTF-8", standalone=True)
    files["word/_rels/fontTable.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    ct = etree.fromstring(files["[Content_Types].xml"])
    CT = "http://schemas.openxmlformats.org/package/2006/content-types"
    if ct.find(f"{{{CT}}}Default[@Extension='odttf']") is None:
        d = etree.SubElement(ct, f"{{{CT}}}Default")
        d.set("Extension", "odttf")
        d.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
        ct.insert(0, d)
    files["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    settings = etree.fromstring(files["word/settings.xml"])
    if settings.find(f"{{{W}}}embedTrueTypeFonts") is None:
        before = ["writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime",
                  "doNotDisplayPageBoundaries", "displayBackgroundShape", "printPostScriptOverText",
                  "printFractionalCharacterWidth", "printFormsData"]
        pos = 0
        for idx, child in enumerate(settings):
            if etree.QName(child).localname in before:
                pos = idx + 1
        settings.insert(pos, etree.Element(f"{{{W}}}embedTrueTypeFonts"))
    files["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist() + [k for k in files if k not in zin.namelist()]:
            zout.writestr(name, files[name])
    return out.getvalue()


# -----------------------------------------------------------------------------
# Document conversion
# -----------------------------------------------------------------------------
class Context:
    def __init__(self, tex_path, bbl_path):
        self.tex_path = os.path.abspath(tex_path)
        self.base = os.path.dirname(self.tex_path)
        self.raw = strip_comments(open(tex_path, encoding="utf-8").read())
        self.japanese = bool(re.search(r"\\documentclass\[[^\]]*\b(ja|japanese)\b", self.raw))
        # Module-level style state belongs to one document: reset it for each conversion
        global JP_MODE
        JP_MODE = self.japanese
        COLORS.clear()
        COLORS.update(DEFAULT_COLORS)
        self.sources = [self.tex_path] + ([os.path.abspath(bbl_path)] if bbl_path and os.path.exists(bbl_path) else [])
        self.bib = parse_bbl(bbl_path)
        self.bbl_path, self.warned_bbl = bbl_path, False
        self.bibcount = 0
        self.icons = load_icon_map()
        self._read_preamble()
        if self.japanese:
            COLORS.update(JP_STYLE["colors"])

    def source_date(self):
        """Newest modification time of the CV's sources (UTC, whole seconds), or
        SOURCE_DATE_EPOCH when set, as reproducible-build tools expect."""
        epoch = os.environ.get("SOURCE_DATE_EPOCH")
        if epoch and epoch.isdigit():
            ts = int(epoch)
        else:
            ts = int(max((os.path.getmtime(p) for p in self.sources if os.path.exists(p)), default=0))
        return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).replace(tzinfo=None)

    def today(self):
        """\\today: the build date like LaTeX, or SOURCE_DATE_EPOCH when set (as LaTeX
        does with FORCE_SOURCE_DATE), so a fixed epoch gives byte-identical output."""
        epoch = os.environ.get("SOURCE_DATE_EPOCH")
        if epoch and epoch.isdigit():
            d = datetime.datetime.fromtimestamp(int(epoch), datetime.timezone.utc).date()
        else:
            d = datetime.date.today()
        if self.japanese:
            return f"{d.year}年{d.month}月{d.day}日"
        return f"{MONTHS[d.month - 1]} {d.day}, {d.year}"

    def space(self, fmt, ems):
        """A space ems wide in the run's font (em/en space glyphs when the font has them)."""
        role = resolve_font(fmt.role, fmt.bold, fmt.italic)[0]
        if ems == 1.0 and has_glyph(role, "\u2003"):
            return "\u2003"
        if ems == 0.5 and has_glyph(role, "\u2002"):
            return "\u2002"
        # No-break spaces, so whitespace collapsing keeps them all
        sw = text_width("\u00a0", fmt.role, fmt.size, fmt.bold, fmt.italic) or fmt.size * 0.25
        return "\u00a0" * max(1, round(ems * fmt.size / sw))

    def _read_preamble(self):
        pre = self.raw.split("\\begin{document}")[0]

        def arg(cmd, n=1):
            m = re.search(r"\\" + cmd + r"(?![A-Za-z])", pre)
            if not m:
                return None
            i, vals = m.end(), []
            for _ in range(n):
                v, i = read_group(pre, i)
                vals.append(v)
            return vals if n > 1 else vals[0]

        self.info = {}
        name = arg("name", 2)
        if name:
            self.info["firstname"], self.info["lastname"] = name
        for cmd in ("firstname", "lastname", "familyname", "position", "address", "mobile", "whatsapp", "email",
                    "dateofbirth", "homepage", "github", "gitlab", "bitbucket", "linkedin", "orcid", "twitter", "x",
                    "researchgate", "skype", "reddit", "xing", "medium", "kaggle", "hackerrank", "telegram",
                    "extrainfo", "quote"):
            v = arg(cmd)
            if v is not None:
                self.info["lastname" if cmd == "familyname" else cmd] = v
        for cmd, n in (("stackoverflow", 2), ("mastodon", 2), ("googlescholar", 2)):
            v = arg(cmd, n)
            if v:
                self.info[cmd] = v
        if re.search(r"\\photo(?![A-Za-z])", pre):
            print("cv2docx: warning: \\photo is not supported in the DOCX; the photo is left out", file=sys.stderr)
        hp = self.info.get("homepage")
        if hp and not hp.startswith(("http://", "https://")):
            self.info["homepage"] = "https://" + hp
        m = re.search(r"\\colorlet\{awesome\}\{([\w-]+)\}", pre)
        if m and m.group(1) in AWESOME_COLORS:
            COLORS["awesome"] = AWESOME_COLORS[m.group(1)]
        m = re.search(r"\\definecolor\{awesome\}\{HTML\}\{([0-9A-Fa-f]{6})\}", pre)
        if m:
            COLORS["awesome"] = m.group(1).upper()
        for cname in ("darktext", "text", "graytext", "lighttext", "sectiondivider"):
            m = re.search(r"\\definecolor\{" + cname + r"\}\{HTML\}\{([0-9A-Fa-f]{6})\}", pre)
            if m:
                COLORS[cname] = m.group(1).upper()
        m = re.search(r"\\setbool\{acvSectionColorHighlight\}\{(\w+)\}", pre)
        self.highlight = not (m and m.group(1) == "false") and not self.japanese
        m = re.search(r"\\renewcommand\{\\acvHeaderSocialSep\}\{(.*)\}\s*$", pre, re.M)
        self.social_sep = m.group(1) if m else "\\quad\\textbar\\quad"
        # Page geometry
        self.paper = (595.28, 841.89) if "letterpaper" not in pre else (612.0, 792.0)
        self.margins = {"left": 2.0 * 28.4528, "right": 2.0 * 28.4528, "top": 1.5 * 28.4528,
                        "bottom": 2.0 * 28.4528, "footskip": 0.5 * 28.4528}
        m = re.search(r"\\geometry\{([^}]*)\}", pre)
        if m:
            for kv in m.group(1).split(","):
                if "=" in kv:
                    k, v = (x.strip() for x in kv.split("=", 1))
                    if k in self.margins:
                        self.margins[k] = parse_length(v if not v.startswith(".") else "0" + v, self.margins[k])
        self.pdfauthor = arg("setpdfauthor") or " ".join(
            x for x in (self.info.get("firstname"), self.info.get("lastname")) if x)
        self.pdftitle = arg("setpdftitle") or "Curriculum Vitae"
        self.body = self.raw.split("\\begin{document}", 1)[-1].split("\\end{document}")[0]

    def read_input(self, name):
        name = name.strip()
        cands = [name, name + ".tex"]
        for base in (self.base, os.getcwd()):
            for c in cands:
                p = os.path.join(base, c)
                if os.path.isfile(p):
                    self.sources.append(os.path.abspath(p))
                    return strip_comments(open(p, encoding="utf-8").read())
        print(f"cv2docx: warning: \\input{{{name}}} not found", file=sys.stderr)
        return ""


def load_icon_map():
    """\\faName -> (role, character) from fontawesome6's mapping file."""
    icons = {}
    sty = None
    try:
        sty = subprocess.run(["kpsewhich", "fontawesome6-mapping.def"], capture_output=True, text=True).stdout.strip()
    except FileNotFoundError:
        pass
    if sty and os.path.exists(sty):
        for m in re.finditer(r"\\__fontawesome_def_icon:nnnnn\{\\(fa\w+)\}\{[^}]*\}\{(free|brands)\d\}\{\d+\}\{\"([0-9A-F]+)\}",
                             open(sty, encoding="utf-8").read()):
            icons[m.group(1)] = ("fasolid" if m.group(2) == "free" else "fabrands", chr(int(m.group(3), 16)))
    return icons


class Converter:
    def __init__(self, ctx):
        self.ctx = ctx
        self.b = Builder(ctx)
        self.inline = Inline(ctx)
        self.refsection = 0

    # --- top level ------------------------------------------------------------------
    def run(self):
        self.walk(self.ctx.body)

    def walk(self, s):
        i = 0
        while i < len(s):
            j = s.find("\\", i)
            if j < 0:
                break
            m = re.match(r"\\([A-Za-z@]+)\*?", s[j:])
            if not m:
                i = j + 2
                continue
            name = m.group(1)
            i = j + len(m.group(0))
            handler = getattr(self, "cmd_" + name, None)
            if handler:
                i = handler(s, i)

    def cmd_input(self, s, i):
        name, i = read_group(s, i)
        self.walk(self.ctx.read_input(name))
        return i

    cmd_include = cmd_input

    def cmd_makecvheader(self, s, i):
        opt, i = read_optional(s, i)
        self.header(opt or "C")
        return i

    def cmd_makecvfooter(self, s, i):
        parts = []
        for _ in range(3):
            v, i = read_group(s, i)
            parts.append(v)
        self.footer(parts)
        return i

    def cmd_pagebreak(self, s, i):
        self.b.page_break()
        return i

    cmd_newpage = cmd_clearpage = cmd_pagebreak

    def cmd_printbibliography(self, s, i):
        """A \\printbibliography outside any refsection lists refsection 0."""
        opts, i = read_optional(s, i)
        self.bibliography(opts, 0)
        return i

    def cmd_end(self, s, i):
        env, i = read_group(s, i)
        if env == "cventries":
            self.b.gap(LAYOUT["entries_after"])
        return i

    def cmd_cvsection(self, s, i):
        title, i = read_group(s, i)
        self.section(title)
        return i

    def cmd_cvsubsection(self, s, i):
        title, i = read_group(s, i)
        self.subsection(self.inline(title, Fmt()))
        return i

    def cmd_cventry(self, s, i):
        args = []
        for _ in range(5):
            v, i = read_group(s, i)
            args.append(v)
        self.entry(*args)
        return i

    def cmd_cvsubentry(self, s, i):
        args = []
        for _ in range(4):
            v, i = read_group(s, i)
            args.append(v)
        self.subentry(*args)
        return i

    def cmd_begin(self, s, i):
        env, i = read_group(s, i)
        if env == "cvhonors":
            loc, i = read_optional(s, i)
            date, i = read_optional(s, i)
            end, after = find_env_end(s, i, env)
            self.honors(s[i:end], loc, date)
            return after
        if env == "cvskills":
            end, after = find_env_end(s, i, env)
            self.skills(s[i:end])
            return after
        if env == "cvparagraph":
            end, after = find_env_end(s, i, env)
            self.paragraph(s[i:end])
            return after
        if env == "refsection":
            end, after = find_env_end(s, i, env)
            self.refsection += 1
            body = s[i:end]
            m = re.search(r"\\printbibliography(?![A-Za-z])", body)
            if m:
                self.bibliography(read_optional(body, m.end())[0], self.refsection)
            return after
        if env == "cvitems":
            end, after = find_env_end(s, i, env)
            self.items(s[i:end], keep_last=False)
            return after
        if env == "cventries":
            self.b.first_in_section = True
        return i

    # --- header & footer -------------------------------------------------------------
    def header(self, align):
        info, b = self.ctx.info, self.b
        al = {"L": WD_ALIGN_PARAGRAPH.LEFT, "R": WD_ALIGN_PARAGRAPH.RIGHT}.get(align, WD_ALIGN_PARAGRAPH.CENTER)
        link = info.get("homepage")
        if self.ctx.japanese:
            size = JP_STYLE["name_size"]
            # The whole name in black, parts joined with "・" (e.g. ジョシ・ラビ・プラカシュ)
            first_f = last_f = Fmt("name", size, "text", bold=True, link=link)
        else:
            first_f = Fmt("headlight", 32, "graytext", link=link)
            last_f = Fmt("head", 32, "text", bold=True, link=link)
        first = self.inline(info.get("firstname", ""), first_f)
        last = self.inline(info.get("lastname", ""), last_f)
        p = b.para(line=LAYOUT["name_line"], align=al, keep=True, before=LAYOUT["page_top_shift"])
        if self.ctx.japanese:
            parts = re.split(r"[\s\u3000\u00a0・]+", plain(first) + " " + plain(last))
            b.add_runs(p, [Seg(JP_STYLE["name_sep"].join(x for x in parts if x), last_f)])
        else:
            b.add_runs(p, first + [Seg(" ", last_f)] + last)
        if info.get("position"):
            p = b.para(line=LAYOUT["position_line"], align=al, keep=True)
            b.add_runs(p, self.inline(info["position"], Fmt("sans", 7.6, "awesome", smallcaps=True)))
        if info.get("address"):
            p = b.para(line=LAYOUT["address_line"], align=al, keep=True)
            b.add_runs(p, self.inline(info["address"], Fmt("head", 8, "lighttext", italic=True)))
        social = self.social_items()
        if social:
            fmt = self.social_fmt()
            sep = self.inline(self.ctx.social_sep, fmt)
            segs = []
            for n, item in enumerate(social):
                if n:
                    segs += sep
                segs += item
            p = b.para(line=LAYOUT["social_line"], align=al, keep=True)
            b.add_runs(p, segs)
        if info.get("quote"):
            p = b.para(line=12, align=al, keep=True)
            b.add_runs(p, self.inline(info["quote"], Fmt("sans", 9, "darktext", italic=True)))
        b.spacer(LAYOUT["after_header"])

    def social_fmt(self):
        return Fmt("head", JP_STYLE["social_size"] if self.ctx.japanese else 6.8, "text")

    def social_items(self):
        info, fmt = self.ctx.info, self.social_fmt()
        items = []

        def add(icon, text, url=None):
            f = replace(fmt, link=url)
            segs = []
            if icon in self.ctx.icons:
                segs.append(self.inline.icon(icon, f))
                segs.append(Seg(" ", f))
            segs += self.inline(text, f)
            items.append(segs)

        if "mobile" in info:
            add("faMobile", info["mobile"], "tel:" + re.sub(r"[ ()\-]", "", info["mobile"]))
        if "whatsapp" in info:
            add("faSquareWhatsapp", info["whatsapp"], "https://wa.me/" + re.sub(r"[ ()+\-]", "", info["whatsapp"]))
        if "email" in info:
            add("faEnvelope", info["email"], "mailto:" + info["email"])
        if "dateofbirth" in info:
            add("faCakeCandles", info["dateofbirth"])
        if "homepage" in info:
            add("faHouseChimney", info["homepage"], info["homepage"])
        simple = [("github", "faSquareGithub", "https://github.com/{}"), ("gitlab", "faGitlab", "https://gitlab.com/{}"),
                  ("bitbucket", "faBitbucket", "https://bitbucket.org/{}")]
        for key, icon, url in simple:
            if key in info:
                add(icon, info[key], url.format(info[key]))
        if "stackoverflow" in info:
            sid, sname = info["stackoverflow"]
            add("faStackOverflow", sname, f"https://stackoverflow.com/users/{sid}")
        more = [("linkedin", "faLinkedin", "https://www.linkedin.com/in/{}"), ("orcid", "faOrcid", "https://orcid.org/{}"),
                ("twitter", "faTwitter", "https://twitter.com/{}"), ("x", "faXTwitter", "https://x.com/{}")]
        for key, icon, url in more:
            if key in info:
                add(icon, info[key], url.format(info[key]))
        if "mastodon" in info:
            inst, mname = info["mastodon"]
            add("faMastodon", mname, f"https://{inst}/@{mname}")
        rest = [("skype", "faSkype", None), ("reddit", "faReddit", "https://www.reddit.com/user/{}"),
                ("researchgate", "faResearchgate", "https://www.researchgate.net/profile/{}"),
                ("xing", "faSquareXing", "https://www.xing.com/profile/{}"), ("medium", "faMedium", "https://medium.com/@{}"),
                ("kaggle", "faKaggle", "https://kaggle.com/{}"), ("hackerrank", "faHackerrank", "https://www.hackerrank.com/{}"),
                ("telegram", "faTelegram", "https://t.me/{}")]
        for key, icon, url in rest:
            if key in info:
                add(icon, info[key], url.format(info[key]) if url else None)
        if "googlescholar" in info:
            gid, gname = info["googlescholar"]
            if not gname.strip():
                gname = f"{info.get('firstname', '')}~{info.get('lastname', '')}"
            add("faGraduationCap", gname, f"https://scholar.google.com/citations?user={gid}")
        if "extrainfo" in info:
            items.append(self.inline(info["extrainfo"], fmt))
        return items

    def footer(self, parts):
        b = self.b
        footer = self.b.doc.sections[0].footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(10)
        tw = b.text_width
        pf.tab_stops.add_tab_stop(Pt(tw / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Pt(tw), WD_TAB_ALIGNMENT.RIGHT)
        fmt = Fmt("sans", 8, "lighttext", smallcaps=True)
        for n, part in enumerate(parts):
            if n:
                b.add_runs(p, [Seg("\t", fmt)])
            b.add_runs(p, self.inline(part, fmt))

    # --- sections ------------------------------------------------------------------------
    def section(self, title):
        b = self.b
        segs = self.inline(title, Fmt("sans", 16, "text", bold=True))
        text = plain(segs)
        b.gap(LAYOUT["section_before"])
        p = b.para(line=LAYOUT["section_line"], after=LAYOUT["section_after"], keep=True, style="Heading 1")
        self._reset_heading(p)
        if self.ctx.highlight and segs:
            hi, rest = text[:3], text[3:]
            runs = [Seg(hi, Fmt("sans", 16, "awesome", bold=True)), Seg(rest, Fmt("sans", 16, "text", bold=True))]
        else:
            runs = [Seg(text, Fmt("sans", 16, "text", bold=True))]
        b.add_runs(p, runs + [Seg(" ", Fmt("sans", 16, "text", bold=True))])
        # The divider is a right tab with a line leader, drawn in the divider colour
        p.paragraph_format.tab_stops.add_tab_stop(Pt(b.text_width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.HEAVY)
        b.add_runs(p, [Seg("\t", Fmt("sans", 9, "sectiondivider"))])
        b.first_in_section = True

    def subsection(self, segs):
        p = self.b.para(line=LAYOUT["subsection_line"], before=LAYOUT["subsection_before"],
                        after=LAYOUT["subsection_after"], keep=True, style="Heading 2")
        self._reset_heading(p)
        self.b.add_runs(p, [Seg(s.text, Fmt("sans", 12, "text", smallcaps=True), s.kind) for s in segs])

    def _reset_heading(self, p):
        """Keep the outline level of Word's heading styles but none of their look."""
        pPr = p._p.get_or_add_pPr()
        for tag in ("w:pBdr", "w:shd"):
            el = pPr.find(qn(tag))
            if el is not None:
                pPr.remove(el)

    def paragraph(self, body):
        p = self.b.para(line=LAYOUT["paragraph_line"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.b.add_runs(p, self.inline(body, Fmt("sanslight", 9, "text")))
        self.b.spacer(5.7, keep=False)

    # --- entries -----------------------------------------------------------------------------
    def entry(self, position, title, location, date, desc):
        b, il = self.b, self.inline
        b.gap(LAYOUT["first_entry_gap"] if b.first_in_section else LAYOUT["entry_gap"])
        b.first_in_section = False
        date_w = 4.5 * 28.4528
        widths = [b.text_width - date_w, date_w]
        has_title = bool((title + location).strip())
        rows = []
        if has_title:
            rows.append((il(title, Fmt("sans", 10, "darktext", bold=True)),
                         il(location, Fmt("sanslight", 9, "awesome", italic=True)), LAYOUT["entry_title_line"]))
        rows.append((il(position, Fmt("sans", 8, "graytext", smallcaps=True)),
                     il(date, Fmt("sanslight", 8, "graytext", italic=True)), LAYOUT["entry_position_line"]))
        t = b.table(widths, len(rows))
        has_desc = bool(desc.strip())
        for r, (left, right, line) in zip(t.rows, rows):
            b.cell_para(r.cells[0], left, line, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
            b.cell_para(r.cells[1], right, line, align=WD_ALIGN_PARAGRAPH.RIGHT,
                        valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
        if not has_desc:
            return
        self.description(desc)

    def description(self, desc):
        if re.search(r"\\cvsubentry(?![A-Za-z])", desc):  # sub-entries (and their own cvitems)
            self.walk(desc)
            return
        m = re.search(r"\\begin\{cvitems\}", desc)
        if m:
            end, after = find_env_end(desc, m.end(), "cvitems")
            pre = desc[:m.start()].strip()
            if pre:
                p = self.b.para(line=LAYOUT["item_line"], keep=True)
                self.b.add_runs(p, self.inline(pre, Fmt("sanslight", 9, "text")))
            self.b.spacer(LAYOUT["items_before"])
            self.items(desc[m.end():end], keep_last=False)
        else:
            p = self.b.para(line=LAYOUT["item_line"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.b.add_runs(p, self.inline(desc, Fmt("sanslight", 9, "text")))

    def items(self, body, keep_last):
        parts = re.split(r"\\item(?![A-Za-z])", body)[1:]
        fmt = Fmt("sanslight", 9, "text")
        for n, part in enumerate(parts):
            keep = keep_last or n < len(parts) - 1
            label, i = read_optional(part, 0)  # \\item[label] text
            text = part[i:].strip()
            self.b.bullet(self.inline(text, fmt), keep,
                          label=None if label is None else self.inline(label, fmt))

    def subentry(self, position, title, date, desc):
        b, il = self.b, self.inline
        date_w = 4.5 * 28.4528
        rows = [(il(title, Fmt("sans", 8, "graytext")), il(date if not position.strip() else "", Fmt("sanslight", 7, "graytext", italic=True)))]
        if position.strip():
            rows.append((il(position, Fmt("sans", 7, "graytext", smallcaps=True)), il(date, Fmt("sanslight", 7, "graytext", italic=True))))
        t = b.table([b.text_width - date_w, date_w], len(rows))
        for r, (left, right) in zip(t.rows, rows):
            b.cell_para(r.cells[0], left, 10, left_pad=5.7)
            b.cell_para(r.cells[1], right, 10, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if desc.strip():
            p = b.para(line=10, left=5.7)
            b.add_runs(p, il(desc, Fmt("sanslight", 8, "text")))

    # --- honors ------------------------------------------------------------------------------
    def honors(self, body, loc_w, date_w):
        b, il = self.b, self.inline
        rows = []
        for m in re.finditer(r"\\cvhonor(?![A-Za-z])", body):
            i = m.end()
            opt, i = read_optional(body, i)
            args = []
            for _ in range(4):
                v, i = read_group(body, i)
                args.append(v)
            rows.append((opt or "", *args))
        if not rows:
            return
        date_w = parse_length(date_w, 1.0 * 28.4528) if date_w is not None else 1.0 * 28.4528
        # Like the class: at least date_w, wider (plus a 1em gap) for long dates such as "2020年"
        date_fmt = Fmt("sans", 9, "graytext")
        widest_date = max(text_width(plain(il(r[4], date_fmt)), "sans", 9) for r in rows)
        if widest_date > date_w:
            date_w = widest_date + 9.0
        loc_fmt = Fmt("sanslight", 9, "awesome", italic=True)
        if loc_w is None or not loc_w.strip():
            widest = max(text_width(plain(il(r[3], loc_fmt)), "sanslight", 9, italic=True) for r in rows)
            loc_w = widest + EM + 1.0
        else:
            loc_w = parse_length(loc_w)
        mid_w = b.text_width - date_w - loc_w
        min_mid = 0.3 * b.text_width
        if mid_w < min_mid:  # an explicit or very long location would squeeze out the text
            loc_w = max(b.text_width - date_w - min_mid, 0.0)
            mid_w = b.text_width - date_w - loc_w
            print(f"cv2docx: warning: cvhonors location column narrowed to {loc_w / 28.4528:.1f}cm "
                  "to leave room for the text", file=sys.stderr)
        widths = [date_w, mid_w] + ([loc_w] if loc_w > 0 else [])
        b.spacer(LAYOUT["honors_before"])
        t = b.table(widths, len(rows))
        for n, (row, (opt, pos, title, loc, date)) in enumerate(zip(t.rows, rows)):
            keep = n < len(rows) - 1
            mid = il(pos, Fmt("sans", 9, "darktext", bold=True))
            if opt not in ("", ":"):
                mid += il(opt, Fmt("sanslight", 8, "text"))
            if title.strip():
                sep = ":" if opt == ":" else ("," if opt == "" else "")
                tf = Fmt("sanslight", 9, "text")
                mid += [Seg(sep + " ", tf)] if sep else [Seg(" ", tf)]
                mid += il(title, tf)
            mid = normalize_space(mid)
            b.cell_para(row.cells[0], il(date, Fmt("sans", 9, "graytext")), LAYOUT["honor_line"], keep=keep)
            b.cell_para(row.cells[1], mid, LAYOUT["honor_line"], keep=keep)
            if loc_w > 0:
                b.cell_para(row.cells[2], il(loc, loc_fmt), LAYOUT["honor_line"], align=WD_ALIGN_PARAGRAPH.RIGHT, keep=keep)
        b.gap(LAYOUT["honors_after"])
        b.first_in_section = False

    # --- skills -------------------------------------------------------------------------------
    def skills(self, body):
        b, il = self.b, self.inline
        rows = []
        for m in re.finditer(r"\\cvskill(?![A-Za-z])", body):
            i = m.end()
            t, i = read_group(body, i)
            v, i = read_group(body, i)
            rows.append((t, v))
        if not rows:
            return
        tf = Fmt("sans", 10, "darktext", bold=True)
        label_w = max(text_width(plain(il(t, tf)), "sans", 10, bold=True) for t, _ in rows) + 2 * EX - 1.5
        b.spacer(LAYOUT["skills_before"])
        t = b.table([label_w, b.text_width - label_w], len(rows))
        for n, (row, (typ, val)) in enumerate(zip(t.rows, rows)):
            keep = n < len(rows) - 1
            b.cell_para(row.cells[0], il(typ, tf), LAYOUT["skill_line"], align=WD_ALIGN_PARAGRAPH.RIGHT,
                        keep=keep, right_pad=EX)
            b.cell_para(row.cells[1], il(val, Fmt("sanslight", 9, "text")), LAYOUT["skill_line"], keep=keep,
                        left_pad=EX, right_pad=EX)
        b.gap(LAYOUT["skills_after"])
        b.first_in_section = False

    # --- bibliography ------------------------------------------------------------------------
    def bibliography(self, opts, number):
        if not self.ctx.bib and not self.ctx.warned_bbl:
            print(f"cv2docx: warning: {self.ctx.bbl_path} not found or empty; publications will be empty",
                  file=sys.stderr)
            self.ctx.warned_bbl = True
        entries = self.ctx.bib.get(number, [])
        self.ctx.bibcount = len(entries)
        title, heading = None, "bibliography"
        if opts:
            tm = re.search(r"title\s*=\s*", opts)
            if tm:
                title, _ = read_group(opts, tm.end())
            hm = re.search(r"heading\s*=\s*([\w-]+)", opts)
            if hm:
                heading = hm.group(1)
        if heading != "none":
            segs = self.inline("References" if title is None else title, Fmt())
            if heading.startswith("subbib"):  # the class's \cvsubsection look
                self.subsection(segs)
            else:  # biblatex's default heading is an unstyled \section*: large bold
                p = self.b.para(line=20, before=10, after=4, keep=True, style="Heading 2")
                self._reset_heading(p)
                self.b.add_runs(p, [Seg(x.text, Fmt("sans", 14.4, "000000", bold=True), x.kind) for x in segs])
        self.b.spacer(LAYOUT["bib_first_gap"])
        for n, e in enumerate(entries):
            if n:
                self.b.gap(LAYOUT["bib_article_gap"] if entries[n - 1]["type"] == "article" else LAYOUT["bib_gap"])
            self.bib_entry(e)
        self.b.gap(LAYOUT["bib_after"])

    def bib_entry(self, e):
        b, il = self.b, self.inline
        f, lists, names = e["fields"], e["lists"], e["names"]
        title_fmt = Fmt("sans", 9, "graytext")
        venue_fmt = Fmt("sanslight", 8, "graytext")
        p = b.para(line=LAYOUT["bib_title_line"], keep=True)
        b.add_runs(p, il(bib_tex(f.get("title", "")), title_fmt))
        authors = ", ".join(format_name(n) for n in names.get("author", []))
        p = b.para(line=LAYOUT["bib_author_line"], keep=True, left=BIBHANG)
        b.add_runs(p, il(bib_tex(authors), title_fmt))
        etype = e["type"]
        year = f.get("year", "")
        ident = self.bib_ident(e, venue_fmt)
        if etype == "article":
            segs = il(bib_tex(f.get("journaltitle", "")), replace(venue_fmt, italic=True))
            vol = f.get("volume", "")
            num = f.get("number", "")
            if vol:
                segs += [Seg(" " + vol + (("." + num) if num else ""), venue_fmt)]
            if year:
                segs += [Seg(f" ({year})", venue_fmt)]
            pages = f.get("pages")
            if pages:
                pp = "pp." if ("bibrangedash" in pages or "-" in pages) else "p."
                segs += [Seg(f" {pp}\u00a0", venue_fmt)] + il(bib_tex(pages), venue_fmt)
            pub = " ".join(lists.get("publisher", []))
            if pub:
                segs += [Seg(". ", venue_fmt)] + il(bib_tex(pub), venue_fmt)
            p = b.para(line=LAYOUT["bib_venue_line"], keep=True, left=BIBHANG)
            b.add_runs(p, normalize_space(segs))
            self.bib_rows([(ident, il(year, venue_fmt))], last_keep=False, line=LAYOUT["bib_article_row_line"])
        elif etype == "inproceedings":
            loc = ", ".join(lists.get("location", []))
            self.bib_rows([(il(bib_tex(f.get("booktitle", "")), replace(venue_fmt, italic=True)),
                            il(bib_tex(loc), Fmt("sanslight", 8, "awesome", italic=True))),
                           (ident, il(year, venue_fmt))], last_keep=False)
        else:
            # incollection/misc: booktitle (italic), [note,] year, location; report:
            # "type number, year" -- the drivers in awesome-cv.cls
            loc = ", ".join(lists.get("location", []))
            parts = []
            if etype == "report":
                tn = " ".join(x for x in (f.get("type", ""), f.get("number", "")) if x)
                parts = [il(bib_tex(tn), venue_fmt) if tn else [], il(year, venue_fmt)]
            else:
                parts = [il(bib_tex(f.get("booktitle", "")), replace(venue_fmt, italic=True))]
                if etype != "incollection":
                    parts.append(il(bib_tex(f.get("note", "")), venue_fmt))
                parts += [il(year, venue_fmt), il(bib_tex(loc), venue_fmt)]
            segs = []
            for part in (x for x in parts if x):  # skip empty fields, so no stray commas
                segs += ([Seg(", ", venue_fmt)] if segs else []) + part
            if segs:
                p = b.para(line=LAYOUT["bib_venue_line"], left=BIBHANG)
                b.add_runs(p, normalize_space(segs))

    def bib_ident(self, e, fmt):
        if "doi" in e["verbs"]:
            label, val, mono = "doi", e["verbs"]["doi"], True
            url = "https://doi.org/" + val
        elif "isbn" in e["fields"]:
            label, val, mono, url = "isbn", e["fields"]["isbn"], False, None
        elif "url" in e["verbs"]:
            label, val, mono = "url", e["verbs"]["url"], True
            url = val
        else:
            return []
        segs = [Seg(label, replace(fmt, smallcaps=True)), Seg(": ", fmt)]
        vf = replace(fmt, role="mono", link=url) if mono else replace(fmt, link=url)
        segs.append(Seg(val, vf))
        return segs

    def bib_rows(self, rows, last_keep, line=None):
        line = line or LAYOUT["bib_row_line"]
        b = self.b
        right_fmt_w = max(text_width(plain(r), "sanslight", 8, italic=True) for _, r in rows) if rows else 0
        right_w = right_fmt_w + EM + 1.0
        total = b.text_width - BIBHANG
        # Indent the table by the hanging indent so it lines up with the authors
        t = b.table([total - right_w, right_w], len(rows))
        tblInd = t._tbl.tblPr.find(qn("w:tblInd"))
        tblInd.set(qn("w:w"), str(int(BIBHANG * 20)))
        for n, (row, (left, right)) in enumerate(zip(t.rows, rows)):
            keep = last_keep or n < len(rows) - 1
            b.cell_para(row.cells[0], left, line, keep=keep, valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
            b.cell_para(row.cells[1], right, line, align=WD_ALIGN_PARAGRAPH.RIGHT, keep=keep,
                        valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("tex", help="top-level .tex file of the CV")
    ap.add_argument("-o", "--output", help="output .docx (default: next to the .tex file)")
    ap.add_argument("--bbl", help="biber .bbl file (default: <name>.bbl next to the .tex file)")
    ap.add_argument("--no-embed", action="store_true", help="do not embed fonts")
    ap.add_argument("--font-cache", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), ".fontcache"),
                    help="directory for converted TrueType fonts")
    args = ap.parse_args(argv)
    stem = os.path.splitext(args.tex)[0]
    out = args.output or stem + ".docx"
    bbl = args.bbl or stem + ".bbl"
    # Fonts and icons come from TeX Live; without them widths would be guesses, no
    # fonts embedded and every icon dropped, so refuse rather than write a poor file
    missing = [f for f in (FONT_FILES["sans"][0], FONT_FILES["fasolid"][0]) if not font_path(f)]
    if missing:
        sys.exit(f"cv2docx: error: TeX Live fonts not found ({', '.join(missing)}); run inside the "
                 "texlive-custom Docker image or another TeX Live with kpsewhich on PATH")
    ctx = Context(args.tex, bbl)
    if re.search(r"\\(makelettertitle|begin\{cvletter\})", ctx.body):
        sys.exit(f"cv2docx: error: {args.tex} is a cover letter; only CVs are supported")
    if not ctx.icons:
        print("cv2docx: warning: fontawesome6-mapping.def not found; icons will be missing", file=sys.stderr)
    conv = Converter(ctx)
    conv.run()
    conv.b.save(out, args.font_cache, embed=not args.no_embed)
    print(f"cv2docx: wrote {out}")


if __name__ == "__main__":
    main()
